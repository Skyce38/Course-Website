"""Offline author check; Python with scikit-learn required, never used by the website.

python scripts/check-ai-knn.py [--source-root <local AI source directory>]
Optional source audit also uses existing pandas/openpyxl; no files are modified.
"""
import argparse
from contextlib import redirect_stdout
from fractions import Fraction as F
import hashlib
import io
from itertools import product
from math import isclose, sqrt
from pathlib import Path
import re

ROOT = Path(__file__).resolve().parents[1]
LESSON = ROOT / "courses/artificial-intelligence/summaries/final-knn.md"
text = LESSON.read_text(encoding="utf-8")
blocks = re.findall(r"^```python\n(.*?)^```", text, re.M | re.S)
assert len(blocks) == 3
namespace, output = {}, io.StringIO()
with redirect_stdout(output):
    for block in blocks:
        exec(compile(block, str(LESSON), "exec"), namespace)
assert output.getvalue().splitlines() == [
    "Drug k=5 drugC", "Drug ranking [8, 10, 4, 3, 2, 7, 1, 6, 9, 5]",
    "Buy k=3 Y", "Buy ranking [1, 4, 8, 2, 5, 9, 10, 3, 6, 7]",
    "sklearn Buy Y", "sizes 105 45", "train accuracy 0.942857",
    "test accuracy 1.000000",
]
neighbors, encode = namespace["neighbors"], namespace["encode"]
drug_x, drug_y = namespace["drug_x"], namespace["drug_y"]
drug_query = namespace["drug_query"]
buy_rows, buy_query = namespace["buy_rows"], namespace["buy_query"]
expected_d2 = [F(v) for v in ["173.010660", "123.001933", "123.000000",
               "65.017698", "625.020488", "197.000500", "170.009049",
               "25.004900", "579.006724", "52.031005"]]
for i, row in enumerate(drug_x):
    exact = sum((F(str(a))-F(str(b)))**2 for a, b in zip(row, drug_query))
    assert exact == expected_d2[i]
    label = drug_y[i]
    assert f"| {i+1} | {sqrt(exact):.6f} | {label} |" in text
assert expected_d2[2] < expected_d2[1]  # do not round before ranking
assert [r[0] for r in namespace["drug_rank"]] == [8,10,4,3,2,7,1,6,9,5]
assert [r[3] for r in namespace["drug_rank"][:5]] == [
    "drugC", "drugY", "drugX", "drugC", "drugC"]
for i, (row, mismatch) in enumerate(zip(buy_rows, [1,2,3,1,2,3,4,1,2,2])):
    assert sum(a != b for a, b in zip(row[:4], buy_query)) == mismatch
    assert sum((a-b)**2 for a, b in zip(encode(row[:4]), encode(buy_query))) == 2*mismatch
    assert f"| {i+1} | {mismatch} | {row[4]} |" in text
assert [r[0] for r in namespace["buy_rank"][:3]] == [1,4,8]
assert neighbors([(0,), (0,), (0,)], ["C","B","A"], (0,), 3)[0] == "A"
assert neighbors([(0,), (0,), (0,)], ["C","B","A"], (0,), 1)[0] == "C"
for call in [lambda: neighbors([], [], (0,), 1),
             lambda: neighbors([(0,)], ["A"], (0,), 0),
             lambda: neighbors([(0,)], ["A"], (0,), 2),
             lambda: neighbors([(0,)], [], (0,), 1),
             lambda: neighbors([(0,)], ["A"], (0,1), 1),
             lambda: encode(("A","Blue","General","Thailand"))]:
    try:
        call()
        raise AssertionError("Invalid input accepted")
    except ValueError:
        pass

# Exhaustive known categorical queries: an independent Hamming-distance oracle.
encoder, model = namespace["encoder"], namespace["model"]
queries = list(product(*namespace["categories"]))
for query in queries:
    expected = sorted((2*sum(a != b for a, b in zip(row[:4], query)), i+1)
                      for i, row in enumerate(buy_rows))
    actual = neighbors(namespace["buy_x"], namespace["buy_y"], encode(query), 3)[1]
    assert [(round(d2), i) for i, d2, _, _ in actual] == expected
    distances, indices = model.kneighbors(encoder.transform([query]), n_neighbors=10)
    observed = sorted((round(d*d), int(i)+1) for d, i in zip(distances[0], indices[0]))
    assert observed == expected
    # Every known query has 9 one-hot values; column permutation preserves distances.
    assert len(encode(query)) == 9
from sklearn.neighbors import KNeighborsClassifier
drug_model = KNeighborsClassifier(n_neighbors=5, algorithm="brute").fit(drug_x, drug_y)
assert drug_model.predict([drug_query])[0] == "drugC"
distances, indices = drug_model.kneighbors([drug_query], n_neighbors=10)
assert (indices[0]+1).tolist() == [8,10,4,3,2,7,1,6,9,5]
assert sum(namespace["y_pred"] == namespace["y_test"]) == 45
assert sum(namespace["knn"].predict(namespace["X_train"]) == namespace["y_train"]) == 99

parser = argparse.ArgumentParser()
parser.add_argument("--source-root", type=Path)
args = parser.parse_args()
if args.source_root:
    import openpyxl
    import pandas as pd
    root = args.source_root
    hashes = {
        "Final/KNN_NB_Classification.pdf": "85ae6bac0e5fe7f98a5e689413bf3e4129aaefc07e8abce23ad8e2865ba33f72",
        "Assignment/input.xlsx": "1a7114d4331b5cc0bbda09f249a6b880489cde69435452818a25126f1cfed2de",
        "Assignment/input_sol.xlsx": "738c8372d7807866fc7f833cf97c6b1717b916fccf9aa1a63187cd54a5f55e07",
        "Assignment/KNN_Example.xlsx": "4ea9038cc3dafddca56359fec892edc8bccd28b2c1efe64f38fccbcd9180766d",
        "Assignment/lab_sol.py": "f8e5737afa13888edc7864658a46eef472f9474278ae3f66a4428b617b93e4b3",
    }
    for file, digest in hashes.items():
        assert hashlib.sha256((root/file).read_bytes()).hexdigest() == digest
    duplicate_count = sum(hashlib.sha256(file.read_bytes()).hexdigest() ==
                          hashes["Assignment/input.xlsx"]
                          for file in (root/"Assignment").glob("*.xlsx"))
    assert duplicate_count == 2  # do not publish private duplicate filename
    raw = openpyxl.load_workbook(root/"Assignment/input.xlsx", data_only=True)
    assert list(raw["Train"].iter_rows(min_row=2, values_only=True)) == buy_rows
    assert tuple(raw["Test"].values)[1][:4] == buy_query
    solved = openpyxl.load_workbook(root/"Assignment/input_sol.xlsx", data_only=False)
    cached = openpyxl.load_workbook(root/"Assignment/input_sol.xlsx", data_only=True)
    for i, row in enumerate(buy_rows, 2):
        assert tuple(solved["KNN"].cell(i,c).value for c in range(7,16)) == encode(row[:4])
        assert solved["KNN"].cell(i,17).value.startswith("=SQRT(")
        expected = sqrt(2*sum(a != b for a, b in zip(row[:4], buy_query)))
        assert isclose(cached["KNN"].cell(i,17).value, expected, abs_tol=1e-12)
    incomplete = [(sheet.title, c.coordinate) for sheet in solved
                  for row in sheet for c in row if c.value == "="]
    assert incomplete and all(sheet != "KNN" for sheet, _ in incomplete)
    drug = openpyxl.load_workbook(root/"Assignment/KNN_Example.xlsx", data_only=True)["KNN"]
    for i, row in enumerate(drug_x, 2):
        assert tuple(drug.cell(i,c).value for c in [1,3,4,5,6,7,8]) == row
        assert drug.cell(i,9).value == drug_y[i-2]
        assert isclose(drug.cell(i,10).value, sqrt(expected_d2[i-2]), abs_tol=1e-12)
    assert tuple(drug.cell(13,c).value for c in [1,3,4,5,6,7,8]) == drug_query
    # Execute only the original KNN segment, substituting the read-only absolute input path.
    lines = (root/"Assignment/lab_sol.py").read_text(encoding="utf-8").splitlines()
    code = "\n".join(lines[:4]+lines[11:43]).replace(
        'pd.read_excel("input.xlsx")', 'pd.read_excel(source_file)')
    original = {"source_file": root/"Assignment/input.xlsx"}
    exec(compile(code, "C01-KNN-only", "exec"), original)
    assert original["result_knn"].tolist() == ["Y"]
    print("PASS: five source hashes, X01/X02 duplicate, raw data/encodings, workbook "
          "distance caches independently checked, incomplete formulas excluded, original C01 KNN segment = Y")
else:
    print("NOT RUN: optional workbook/source audit (supply --source-root)")
print("PASS: 3 Python blocks, exact Drug/Buy distances/rankings/votes, 24 categorical "
      "queries against Hamming/sklearn, ties/invalid inputs; Iris train 99/105, test 45/45")
