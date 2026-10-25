"""Offline Stage C checks: independent vectorized references; original sources optional read-only."""
from pathlib import Path
from contextlib import redirect_stdout
from fractions import Fraction as F
import argparse,io,re
import numpy as np
ROOT=Path(__file__).resolve().parents[1]
def load(slug):
 text=(ROOT/f'courses/artificial-intelligence/summaries/{slug}.md').read_text(encoding='utf8')
 blocks=re.findall(r'^```python\n(.*?)^```',text,re.M|re.S);assert len(blocks)==1
 ns={}
 with redirect_stdout(io.StringIO()):exec(compile(blocks[0],slug,'exec'),ns)
 return ns,text
k,kt=load('final-kmeans');p,pt=load('final-perceptron')
km=k['kmeans_trace']; train=p['train_perceptron']
def oracle(points,centers,cap=100):
 a=np.array(points,float);c=np.array(centers,float);trace=[]
 for _ in range(cap):
  d=((a[:,None]-c)**2).sum(axis=2);labels=d.argmin(axis=1)
  new=np.array([a[labels==j].mean(axis=0) if (labels==j).any() else c[j] for j in range(len(c))])
  trace.append((c.copy(),d,labels,new,((a-new[labels])**2).sum()))
  if np.allclose(new,c,atol=1e-13,rtol=0):return trace
  c=new
 raise AssertionError('Cap')
def compare(points,centers):
 actual=km(points,centers);ref=oracle(points,centers);assert len(actual)==len(ref)
 for (_,old,d,labels,new,sse),(oc,od,ol,on,os) in zip(actual,ref):
  assert np.allclose(np.array(old,float),oc) and np.allclose(np.array(d,float),od)
  assert labels==ol.tolist() and np.allclose(np.array(new,float),on) and np.isclose(float(sse),os)
 assert all(actual[i][-1]>=actual[i+1][-1] for i in range(len(actual)-1))
 return actual
x=compare(k['workbook_points'],[(2,10),(5,8),(1,2)]);a=compare(k['customers'],k['customers'][:3])
assert [v[-1] for v in x]==[F(37),F(93,4),F(43,3),F(43,3)]
assert a[-1][-1]==F(183725,6) and len(a)==3
assert a[1][2][7][0]==F('5261.328125') and '5261.328125' in kt
rng=np.random.default_rng(641)
for _ in range(40):
 data=rng.integers(-8,9,size=(12,2)).tolist();compare(data,data[:3])
compare([[0,0],[2,0]],[[1,0],[1,0],[100,0]]) # tie and empty clusters
for bad in [([],[[0,0]]),([[1,2]],[]),([[1]],[[1,2]])]:
 try:km(*bad);raise AssertionError('Bad input')
 except ValueError:pass
tr=train(p['patterns_and'],[.1,.1,.1],.5,6)
assert len(tr)==24 and tr[-1][-1]==(F('-1.4'),F('1.1'),F('.6'))
assert [int(tr[-1][-1][0]+tr[-1][-1][1]*a+tr[-1][-1][2]*b>0) for a,b,_ in p['patterns_and']]==[0,0,0,1]
assert train([(0,0,1)],[0,0,0],.5,1)[0][4:6]==(0,1)
# Independent array update: compare every sample under varied initial values/rates
for _ in range(30):
 initial=rng.integers(-6,7,3)/10;rate=float(rng.choice([.1,.5,.7]));trace=train(p['patterns_and'],initial,rate,7);w=initial.copy()
 for epoch,i,before,z,y,e,after in trace:
  aa,bb,t=p['patterns_and'][i-1];v=np.array([1,aa,bb]);score=w@v
  # Exact decimal reference for threshold around zero; floats alone may round across 0
  exact=sum(F(str(q))*int(vv) for q,vv in zip(before,v));pred=int(exact>0)
  assert np.isclose(float(z),score) and pred==y
  w+=rate*(t-pred)*v
  assert np.allclose(np.array(after,float),w,atol=1e-12)
for rate,epochs in [(0,1),(-1,1),(.5,0)]:
 try:train(p['patterns_and'],[0,0,0],rate,epochs);raise AssertionError('Invalid training setting')
 except ValueError:pass
assert len(train(p['patterns_xor'],[.1,.1,.1],.5,20))==80
args=argparse.ArgumentParser();args.add_argument('--source-root',type=Path);args=args.parse_args()
if args.source_root:
 from openpyxl import load_workbook
 from pypdf import PdfReader
 text=PdfReader(args.source_root/'Assignment/Assignment_KMeans.pdf').pages[0].extract_text()
 rows=re.findall(r'P\d\d\s+(\d+)\s+(\d+)',text)
 assert [(int(a),int(b)) for a,b in rows]==k['customers']
 ws=load_workbook(args.source_root/'Assignment/Demo_kmeans.xlsx',data_only=True).active
 for it,trace in enumerate(x):
  _,old,d,labels,new,sse=trace;start=[6,20,34,48][it]
  for i in range(8):
   assert tuple(ws.cell(start+i,j).value for j in [1,2])==k['workbook_points'][i]
   assert all(np.isclose(float(d[i][j])**.5,ws.cell(start+i,3+j).value) for j in range(3))
   assert ws.cell(start+i,6).value==f'C{labels[i]+1}'
  cstart=[15,29,43,None][it]
  if cstart:
   assert np.allclose(np.array(new,float),[[ws.cell(cstart+j,col).value for col in [2,3]] for j in range(3)])
 ws=load_workbook(args.source_root/'Assignment/perceptron.xlsx',data_only=True).active
 for r,(_,i,before,z,y,e,after) in enumerate(tr,4):
  assert tuple(ws.cell(r,j).value for j in range(1,4))==p['patterns_and'][i-1]
  assert np.allclose([float(z),y,e]+list(map(float,after)),[ws.cell(r,j).value for j in range(4,10)])
 print('PASS: A02 all15 points; X05 all96 distances/32 assignments/centroid rounds; X06 all24 sequential source rows')
print('PASS: exact lesson traces, 40 independent K-Means cases, 30 perceptron references, ties/empty/threshold/XOR boundaries')
