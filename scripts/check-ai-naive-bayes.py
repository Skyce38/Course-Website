"""Offline author check. Requires scikit-learn; never runs in the website.

python scripts/check-ai-naive-bayes.py [--source-root <local AI source directory>]
Optional workbook/source audit uses pandas/openpyxl, reads only and writes nothing.
"""
import argparse
from contextlib import redirect_stdout
from fractions import Fraction as F
import hashlib
import io
from itertools import product
from math import exp, fsum, isclose, log, pi, prod, sqrt
from pathlib import Path
import re

ROOT = Path(__file__).resolve().parents[1]
LESSON = ROOT / "courses/artificial-intelligence/summaries/final-naive-bayes.md"
text = LESSON.read_text(encoding="utf-8")
blocks = re.findall(r"^```python\n(.*?)^```", text, re.M | re.S)
assert len(blocks) == 3
ns, output = {}, io.StringIO()
with redirect_stdout(output):
    for block in blocks:
        exec(compile(block, str(LESSON), "exec"), ns)
scores, normalize = ns["categorical_scores"], ns["normalize"]
assert ns["judy"] == {"+": F(1,18), "-": F(1,125)}
assert normalize(ns["judy"]) == {"+": F(125,143), "-": F(18,143)}
assert sum(ns["judy"].values()) == F(143,2250)
assert ns["judy_labels"].count("+") == 3
assert ns["judy_labels"].count("-") == 5
assert ns["judy_query"] == ("Blonde","Average","Heavy","No")
assert F(2,3)*F(3,8)/F(4,8) == F(1,2)

buy_rows, buy_y, query, categories = (ns[k] for k in [
    "buy_rows", "buy_labels", "buy_query", "buy_categories"])
raw = scores(buy_rows, buy_y, query, categories)
smooth = scores(buy_rows, buy_y, query, categories, 1)
assert raw == {"N": F(0), "Y": F(32,625)}
assert smooth == {"N": F(15,2744), "Y": F(225,5488)}
assert normalize(raw) == {"N": F(0), "Y": F(1)}
assert normalize(smooth) == {"N": F(2,17), "Y": F(15,17)}

# Validate both published full frequency tables, not only query products.
for rows, labels, names, classes in [
    (ns["judy_rows"],ns["judy_labels"],["Hair","Height","Weight","Lotion"],["+","-"]),
    (buy_rows,buy_y,["Model","Color","Type","Origin"],["Y","N"]),
]:
    for j, name in enumerate(names):
        for value in sorted({row[j] for row in rows}):
            counts = [sum(row[j] == value and y == c for row, y in zip(rows, labels))
                      for c in classes]
            assert f"| {name}={value} | {counts[0]} | {counts[1]} |" in text

# Independent encoded-category frequency oracle through sklearn across all known queries.
from sklearn.naive_bayes import CategoricalNB
from sklearn.preprocessing import OrdinalEncoder
for name, rows, labels, vocab in [
    ("Buy",buy_rows,buy_y,categories),
    ("Judy",ns["judy_rows"],ns["judy_labels"],ns["judy_categories"]),
]:
    encoder = OrdinalEncoder(categories=[list(v) for v in vocab])
    encoded = encoder.fit_transform(rows).astype(int)
    model = CategoricalNB(alpha=1.0, force_alpha=True).fit(encoded, labels)
    evidence_total = F(0)
    queries = list(product(*vocab))
    for q in queries:
        result = scores(rows, labels, q, vocab, 1)
        posterior = normalize(result)
        assert sum(posterior.values()) == 1 and all(v >= 0 for v in posterior.values())
        evidence_total += sum(result.values())
        observed = model.predict_proba(encoder.transform([q]).astype(int))[0]
        for c, value in zip(model.classes_, observed):
            assert isclose(value, float(posterior[c]), abs_tol=1e-12)
    assert evidence_total == 1
    assert scores(rows, labels, tuple(queries[0]), vocab, 1) == scores(
        rows[::-1],labels[::-1],tuple(queries[0]),vocab,1)

zero_rows, zero_y, zero_vocab = [("Red",),("Blue",)],["A","B"],[("Red","Blue","Green")]
zero = scores(zero_rows,zero_y,("Green",),zero_vocab)
try:
    normalize(zero)
    raise AssertionError("Zero evidence must fail")
except ValueError:
    pass
assert normalize(scores(zero_rows,zero_y,("Green",),zero_vocab,1)) == {"A":F(1,2),"B":F(1,2)}
for call in [lambda: scores([],[],query,categories),
             lambda: scores(buy_rows,buy_y,query,categories,-1),
             lambda: scores(buy_rows,buy_y,query,categories,.5),
             lambda: scores(buy_rows,buy_y,("C","Red","General","Thailand"),categories),
             lambda: ns["gaussian_density"](0,0,0)]:
    try:
        call()
        raise AssertionError("Invalid input accepted")
    except ValueError:
        pass
assert ns["query_encoded"].tolist() == [[0,0,0,2]]
assert ns["model"].n_categories_.tolist() == [2,2,2,3]
assert ns["model"].predict(ns["query_encoded"]).tolist() == ["Y"]
assert isclose(ns["score_a"], .5/sqrt(2*pi), abs_tol=1e-12)
assert isclose(ns["score_b"], .5*exp(-2)/sqrt(2*pi), abs_tol=1e-12)
assert isclose(ns["score_a"]/(ns["score_a"]+ns["score_b"]),1/(1+exp(-2)),abs_tol=1e-12)
assert sum(ns["y_pred"] == ns["y_test"]) == 44
assert sum(ns["gnb"].predict(ns["X_train"]) == ns["y_train"]) == 99

# Independent Gaussian likelihood calculation, including source-model variance smoothing.
train, targets = ns["X_train"], ns["y_train"]
global_means = [fsum(float(row[j]) for row in train)/len(train) for j in range(4)]
epsilon = 1e-9 * max(fsum((float(row[j])-global_means[j])**2 for row in train)/len(train)
                     for j in range(4))
parameters = []
for c in (0,1,2):
    group = [row for row,y in zip(train,targets) if y == c]
    means = [fsum(float(row[j]) for row in group)/len(group) for j in range(4)]
    variances = [fsum((float(row[j])-means[j])**2 for row in group)/len(group)+epsilon
                 for j in range(4)]
    parameters.append((len(group)/len(train),means,variances))
for row,expected in zip(ns["X_test"],ns["y_pred"]):
    logs = [log(prior)+fsum(-.5*log(2*pi*v)-.5*(float(x)-m)**2/v
                           for x,m,v in zip(row,means,variances))
            for prior,means,variances in parameters]
    assert max(range(3),key=lambda c:logs[c]) == int(expected)
    weights = [exp(v-max(logs)) for v in logs]
    observed = ns["gnb"].predict_proba([row])[0]
    for a,b in zip(weights,observed):
        assert isclose(a/sum(weights),float(b),abs_tol=1e-10)

parser = argparse.ArgumentParser()
parser.add_argument("--source-root", type=Path)
args = parser.parse_args()
if args.source_root:
    import openpyxl
    root = args.source_root
    hashes = {
        "Final/KNN_NB_Classification.pdf":"85ae6bac0e5fe7f98a5e689413bf3e4129aaefc07e8abce23ad8e2865ba33f72",
        "Assignment/input.xlsx":"1a7114d4331b5cc0bbda09f249a6b880489cde69435452818a25126f1cfed2de",
        "Assignment/input_sol.xlsx":"738c8372d7807866fc7f833cf97c6b1717b916fccf9aa1a63187cd54a5f55e07",
        "Assignment/lab_sol.py":"f8e5737afa13888edc7864658a46eef472f9474278ae3f66a4428b617b93e4b3",
    }
    for file,digest in hashes.items():
        assert hashlib.sha256((root/file).read_bytes()).hexdigest() == digest
    source = openpyxl.load_workbook(root/"Assignment/input.xlsx",data_only=True)
    raw_rows = list(source["Train"].values)[1:]
    assert [r[:4] for r in raw_rows] == buy_rows
    assert [r[4] for r in raw_rows] == buy_y
    assert list(source["Test"].values)[1][:4] == query
    formula = openpyxl.load_workbook(root/"Assignment/input_sol.xlsx",data_only=False)["NaiveBayes"]
    cached = openpyxl.load_workbook(root/"Assignment/input_sol.xlsx",data_only=True)["NaiveBayes"]
    assert formula["L3"].value == formula["L6"].value == "="
    assert formula["M3"].value == "=G3*H3*I3*J3*K3"
    assert formula["M6"].value == "=G6*H6*I6*J6*K6"
    for row,expected in [(3,[F(1,2),F(4,5),F(4,5),F(2,5),F(2,5)]),
                         (6,[F(1,2),F(2,5),F(0),F(4,5),F(1,5)])]:
        for col,value in zip("GHIJK",expected):
            assert isclose(cached[f"{col}{row}"].value,float(value),abs_tol=1e-12)
        assert isclose(cached[f"M{row}"].value,float(prod(expected)),abs_tol=1e-12)
    # Execute the original NB segment only, with inputs read from source and path-free data.
    import pandas as pd
    lines=(root/"Assignment/lab_sol.py").read_text(encoding="utf-8").splitlines()
    original={"X":pd.DataFrame(buy_rows,columns=["Model","Color","Type","Origin"]),
              "y":pd.Series(buy_y),"test":pd.DataFrame([query],columns=["Model","Color","Type","Origin"]),
              "OrdinalEncoder":OrdinalEncoder,"CategoricalNB":CategoricalNB}
    exec(compile("\n".join(lines[49:59]),"C01-NB-only","exec"),original)
    assert original["result_nb"].tolist() == ["Y"]
    assert original["nb"].get_params()["alpha"] == 1.0
    assert isclose(original["nb"].predict_proba(original["test2"])[0,1],15/17,abs_tol=1e-12)
    print("PASS: four source fingerprints/raw Buy/worksheet formula-cache audit; "
          "L3/L6 excluded; original C01 NB-only result Y, alpha=1")
else:
    print("NOT RUN: optional local-source audit (supply --source-root)")
print("PASS: 3 lesson blocks, 2 full frequency tables, exact Judy/Buy scores and "
      "normalization, 78 categorical queries vs sklearn, zero/vocabulary cases, "
      "independent Gaussian scores on 45 test rows; Iris 99/105 train and 44/45 test")
