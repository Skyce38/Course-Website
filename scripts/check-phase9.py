"""Independent Phase 9 counterexamples and correction guards; no source writes."""
from fractions import Fraction as F
from pathlib import Path
import argparse
import math
import re
import io
import json
from itertools import product
from contextlib import redirect_stdout

ROOT = Path(__file__).resolve().parents[1]

# A degenerate linear classifier is constant, not a vertical separating line.
points = [(F(x), F(y)) for x in (-10, 0, 7) for y in (-8, 0, 9)]
for bias, expected in [(F(-1), 0), (F(0), 0), (F(1), 1)]:
    assert {int(bias + F(0)*x + F(0)*y > 0) for x, y in points} == {expected}
for y in (-10, 0, 10):
    assert F(-2) + F(4)*F(1, 2) + F(0)*y == 0
text = (ROOT/'courses/artificial-intelligence/summaries/final-perceptron.md').read_text(encoding='utf8')
assert 'w2=0 แต่ w1≠0' in text and 'w1=w2=0' in text

# Workbook R8 differs from the slide's sex field. Exact decimal arithmetic
# isolates the additional two squared one-hot differences, without changing data.
query = (F(36), F(1), F(0), F(1), F(3), F('.697'), F('.069'))
workbook = (F(41), F(1), F(0), F(1), F(3), F('.767'), F('.069'))
slide = (F(41), F(0), F(1), F(1), F(3), F('.767'), F('.069'))
sq = lambda row: sum((a-b)**2 for a, b in zip(row, query))
assert sq(workbook) == F('25.0049') and sq(slide) == F('27.0049')
assert sq(slide)-sq(workbook) == 2
assert math.isclose(math.sqrt(sq(slide)), 5.196623904036158, abs_tol=1e-12)
knn = (ROOT/'courses/artificial-intelligence/summaries/final-knn.md').read_text(encoding='utf8')
assert 'KNN!B9:D9' in knn and '27.0049' in knn and 'ข้อมูล X04' in knn
nb = (ROOT/'courses/artificial-intelligence/summaries/final-naive-bayes.md').read_text(encoding='utf8')
assert 'Paola  Blonde  Tall' in nb and 'Dana   Blonde  Tall' not in nb

# Rebuild uses duplicate-checking insert, so collisions make its comparison
# count triangular rather than linear. Count operations, not noisy timings.
hash_source = (ROOT/'courses/oop-python/summaries/hash-tables.md').read_text(encoding='utf8')
ns = {}
with redirect_stdout(io.StringIO()):
    exec(compile(re.search(r'^```python\n(.*?)^```', hash_source, re.M|re.S)[1], 'hash-lesson', 'exec'), ns)
class Collision:
    comparisons = 0
    def __init__(self, value): self.value = value
    def __hash__(self): return 0
    def __eq__(self, other):
        Collision.comparisons += 1
        return isinstance(other, Collision) and self.value == other.value
for capacity in (8, 16, 32, 64):
    count = 3*capacity//4
    table = ns['ChainingHashTable'](capacity)
    keys = [Collision(i) for i in range(count+1)]
    for i in range(count): table.insert(keys[i], i)
    Collision.comparisons = 0
    table.insert(keys[-1], count)
    assert Collision.comparisons == count*(count+1)//2
    assert len(table) == count+1 and table._capacity == 2*capacity
    assert [table.get(key) for key in keys] == list(range(count+1))
assert 'O(n²+capacity)' in hash_source and 'คาดหมาย O(n+capacity)' in hash_source
array = (ROOT/'courses/oop-python/summaries/arrays-dynamic-array.md').read_text(encoding='utf8')
assert '|O(n−i+1)' in array
print('PASS: four collision-heavy rebuild operation counts and end-insertion complexity qualification')

# Three equally weighted records cannot give empirical entropy exactly one bit,
# even when all three possible labels are allowed. Enumerate every labelling.
entropy = lambda counts: -sum((n/sum(counts))*math.log2(n/sum(counts)) for n in counts if n)
for labels in product(range(3), repeat=3):
    assert not math.isclose(entropy([labels.count(c) for c in range(3)]), 1, abs_tol=1e-12)
assert (2*entropy([2, 0])+6*entropy([3, 3]))/8 == .75
bank = lambda course, slug: json.loads((ROOT/f'courses/{course}/exam/{slug}.json').read_text(encoding='utf8'))
sorting = next(q for q in bank('oop-python','complexity-basic-sorting')['questions'] if q['id'].endswith('04-r2'))
assert 'แคบที่สุด' in sorting['prompt'] and sorting['answer'] == 'a'
for n in range(1, 1001): assert n <= n*n  # Both were upper bounds; the prompt must disambiguate.
tree = next(q for q in bank('artificial-intelligence','final-decision-tree')['questions'] if q['id'].endswith('19-r2'))
assert '2 และ 6 records' in tree['prompt'] and tree['answer'] == .75
intro = next(q for q in bank('artificial-intelligence','ai-foundations-agents')['questions'] if q['id'].endswith('08-r2'))
assert 'ChatGPT' in intro['prompt'] and 'กฎหมาย' not in intro['prompt'] and intro['answer'] == 'b'
print('PASS: exhaustive 3-record entropy counterexample and realizable replacement, tightest-bound and source-premise correction guards')

class Example:
    shared = []
    def method(self): return 7
assert isinstance(Example, type) and isinstance(Example, object)
assert callable(getattr(Example(), 'method'))
a,b = Example(),Example()
a.shared.append(3)
assert b.shared == [3] and a.shared is b.shared
prose = (ROOT/'courses/oop-python/summaries/classes-objects.md').read_text(encoding='utf8')
assert 'ตัว class เองก็เป็น object' in prose

# Execute the original two Linked List blocks in their documented shared scope.
# Check reachable values/identity against plain-list expectations, and compare
# the declarative traversal's output to the executed print stream.
linked = (ROOT/'courses/oop-python/summaries/linked-list.md').read_text(encoding='utf8')
blocks = re.findall(r'^```python\n(.*?)^```', linked, re.M|re.S)
assert len(blocks) == 2
scope, output = {}, io.StringIO()
with redirect_stdout(output): exec(compile(blocks[0], 'linked-traversal', 'exec'), scope)
assert output.getvalue().splitlines() == ['10', '20', '30']
assert scope['current'] is None and scope['head'].data == 10 and scope['tail'].data == 30
original_tail = scope['tail']
exec(compile(blocks[1], 'linked-insertion', 'exec'), scope)
node, seen, values = scope['head'], set(), []
while node is not None:
    assert id(node) not in seen
    seen.add(id(node)); values.append(node.data); node = node.next
assert values == [10, 20, 25, 30] and scope['tail'] is original_tail
stepper = json.loads(re.search(r'^```study-stepper\n(.*?)^```', linked, re.M|re.S)[1])
assert len(stepper['steps']) == 11
assert stepper['steps'][-1]['output'].splitlines() == output.getvalue().splitlines()
intro_text = (ROOT/'courses/artificial-intelligence/summaries/ai-foundations-agents.md').read_text(encoding='utf8')
assert 'ไม่เกิน 70%' in intro_text and '**1966**' in intro_text and 'Simple Reflex' in intro_text
print('PASS: class/object counterexample; both original Linked List blocks and 11-state final output; primary-reference qualification guards')

parser = argparse.ArgumentParser()
parser.add_argument('--source-root', type=Path)
args = parser.parse_args()
if args.source_root:
    from openpyxl import load_workbook
    book = load_workbook(args.source_root/'Assignment/KNN_Example.xlsx', read_only=True, data_only=False)
    assert tuple(book['KNN'].cell(9, c).value for c in (1, 2, 3, 4)) == (41, 'F', 1, 0)
    book.close()
    print('PASS: read-only original X04 R8 identity and encoding')
else:
    print('NOT RUN: optional local-source cell check (--source-root)')
print('PASS: degenerate boundaries, exact R8 counterfactual distance, source-name and visible correction guards')
